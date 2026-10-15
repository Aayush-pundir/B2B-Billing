import datetime
import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Account, JournalEntry, JournalLine, User
from app.routers.core_routes import log_action
from app.routers.purchases_routes import assert_open
from app.services import finreports, ledger, settings_store
from app.services.auth import FINANCE_READERS, require_roles, require_write
from app.services.chart import ACCOUNT_TYPES

router = APIRouter(prefix="/api/finance")
READ = require_roles(*FINANCE_READERS)


def _d(value: str | None, name: str, default: str) -> str:
    try:
        return datetime.date.fromisoformat(value or default).isoformat()
    except ValueError:
        raise HTTPException(400, f"{name} must be a date (YYYY-MM-DD)")


def _fy_start(today: datetime.date) -> str:
    year = today.year if today.month >= 4 else today.year - 1  # Indian financial year starts 1 April
    return datetime.date(year, 4, 1).isoformat()


# ----------------------------- chart of accounts -----------------------------

class AccountIn(BaseModel):
    code: str
    name: str
    type: str
    description: str = ""
    active: bool = True


@router.get("/accounts")
def list_accounts(db: Session = Depends(get_db), user: User = Depends(READ)):
    bal = {}
    for r in ledger.entries(db):
        b = bal.setdefault(r["account"], [0.0, 0.0])
        b[0] += r["debit"]; b[1] += r["credit"]
    out = []
    for a in db.query(Account).order_by(Account.code).all():
        d, c = bal.get(a.code, [0.0, 0.0])
        out.append({"id": a.id, "code": a.code, "name": a.name, "type": a.type, "system": a.system, "active": a.active,
                    "description": a.description, "balance": ledger.signed(a.type, d, c)})
    return out


@router.post("/accounts")
def create_account(p: AccountIn, db: Session = Depends(get_db), user: User = Depends(require_write("accounting"))):
    code = p.code.strip()
    if not code.isdigit() or not 3 <= len(code) <= 8:
        raise HTTPException(400, "Account code must be 3–8 digits")
    if p.type not in ACCOUNT_TYPES:
        raise HTTPException(400, f"Type must be one of {ACCOUNT_TYPES}")
    if not p.name.strip():
        raise HTTPException(400, "Name is required")
    if db.query(Account).filter(Account.code == code).first():
        raise HTTPException(400, "That account code already exists")
    db.add(Account(code=code, name=p.name.strip(), type=p.type, description=p.description))
    db.commit()
    log_action(db, user, "account.create", "account", f"{code} {p.name}")
    return {"ok": True}


@router.put("/accounts/{account_id}")
def update_account(account_id: int, p: AccountIn, db: Session = Depends(get_db), user: User = Depends(require_write("accounting"))):
    a = db.query(Account).get(account_id)
    if not a:
        raise HTTPException(404, "Account not found")
    if a.system and (p.type != a.type or p.code != a.code or not p.active):
        raise HTTPException(400, "System accounts can be renamed but not re-typed, re-coded or deactivated")
    if a.type != p.type and any(r["account"] == a.code for r in ledger.entries(db)):
        raise HTTPException(400, "This account has transactions — its type cannot change")
    a.name, a.description, a.active, a.type = p.name.strip() or a.name, p.description, p.active, p.type
    db.commit()
    return {"ok": True}


# ----------------------------- journals -----------------------------

class JLine(BaseModel):
    account_code: str
    debit: float = 0.0
    credit: float = 0.0
    note: str = ""


class JournalIn(BaseModel):
    date: str
    memo: str = ""
    ref: str = ""
    kind: str = "manual"  # manual | opening
    lines: list[JLine]


@router.get("/journals")
def list_journals(db: Session = Depends(get_db), user: User = Depends(READ)):
    names = {a.code: a.name for a in db.query(Account).all()}
    return [{"id": j.id, "date": j.date, "memo": j.memo, "ref": j.ref or f"JE-{j.id}", "kind": j.kind, "reverses_id": j.reverses_id,
             "created_by": j.created_by, "reversed": any(x.reverses_id == j.id for x in db.query(JournalEntry).filter(JournalEntry.reverses_id == j.id)),
             "lines": [{"account_code": l.account_code, "account": names.get(l.account_code, "?"), "debit": l.debit, "credit": l.credit, "note": l.note} for l in j.lines]}
            for j in db.query(JournalEntry).order_by(JournalEntry.date.desc(), JournalEntry.id.desc()).limit(200).all()]


def _validate_lines(db: Session, lines: list[JLine]):
    if len(lines) < 2:
        raise HTTPException(400, "A journal needs at least two lines")
    accts = {a.code: a for a in db.query(Account).all()}
    for n, l in enumerate(lines, start=1):
        a = accts.get(l.account_code)
        if not a or not a.active:
            raise HTTPException(400, f"Line {n}: account {l.account_code} does not exist or is inactive")
        if l.debit < 0 or l.credit < 0 or (l.debit and l.credit) or not (l.debit or l.credit):
            raise HTTPException(400, f"Line {n}: enter either a debit or a credit amount")
    dr, cr = round(sum(l.debit for l in lines), 2), round(sum(l.credit for l in lines), 2)
    if abs(dr - cr) > 0.005:
        raise HTTPException(400, f"Journal does not balance: debits ₹{dr:,.2f} vs credits ₹{cr:,.2f}")


@router.post("/journals")
def create_journal(p: JournalIn, db: Session = Depends(get_db), user: User = Depends(require_write("accounting"))):
    date = _d(p.date, "Date", "")
    assert_open(db, date)
    if p.kind not in ("manual", "opening"):
        raise HTTPException(400, "Invalid journal kind")
    _validate_lines(db, p.lines)
    je = JournalEntry(date=date, memo=p.memo, ref=p.ref, kind=p.kind, created_by=user.email)
    je.lines = [JournalLine(account_code=l.account_code, debit=l.debit, credit=l.credit, note=l.note) for l in p.lines]
    db.add(je)
    db.commit()
    log_action(db, user, "journal.create", "journal", f"JE-{je.id} {p.memo}")
    return {"id": je.id}


@router.post("/journals/{journal_id}/reverse")
def reverse_journal(journal_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("accounting"))):
    j = db.query(JournalEntry).get(journal_id)
    if not j:
        raise HTTPException(404, "Journal not found")
    if j.kind == "reversal" or db.query(JournalEntry).filter(JournalEntry.reverses_id == j.id).first():
        raise HTTPException(400, "This journal is already reversed (or is itself a reversal)")
    today = datetime.date.today().isoformat()
    assert_open(db, today)
    rev = JournalEntry(date=today, memo=f"Reversal of JE-{j.id}: {j.memo}", ref=f"REV-JE-{j.id}", kind="reversal", reverses_id=j.id, created_by=user.email)
    rev.lines = [JournalLine(account_code=l.account_code, debit=l.credit, credit=l.debit, note=l.note) for l in j.lines]
    db.add(rev)
    db.commit()
    log_action(db, user, "journal.reverse", "journal", f"JE-{j.id}")
    return {"id": rev.id}


# ----------------------------- period lock -----------------------------

class LockIn(BaseModel):
    date: str


@router.get("/lock")
def get_lock(db: Session = Depends(get_db), user: User = Depends(READ)):
    return {"lock_date": settings_store.get_lock_date(db)}


@router.put("/lock")
def set_lock(p: LockIn, db: Session = Depends(get_db), user: User = Depends(require_write("accounting"))):
    value = _d(p.date, "Lock date", "") if p.date else ""
    settings_store.set_lock_date(db, value)
    log_action(db, user, "books.lock", "accounting", value or "unlocked")
    return {"lock_date": value}


# ----------------------------- reports -----------------------------

@router.get("/ledger")
def general_ledger(account: str, start: str | None = None, end: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    today = datetime.date.today()
    res = finreports.general_ledger(db, account, _d(start, "Start", _fy_start(today)), _d(end, "End", today.isoformat()))
    if "error" in res:
        raise HTTPException(404, res["error"])
    return res


@router.get("/reports/trial-balance")
def trial_balance(as_of: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    return finreports.trial_balance(db, _d(as_of, "As-of date", datetime.date.today().isoformat()))


@router.get("/reports/profit-loss")
def profit_loss(start: str | None = None, end: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    today = datetime.date.today()
    return finreports.profit_and_loss(db, _d(start, "Start", _fy_start(today)), _d(end, "End", today.isoformat()))


@router.get("/reports/balance-sheet")
def balance_sheet(as_of: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    return finreports.balance_sheet(db, _d(as_of, "As-of date", datetime.date.today().isoformat()))


@router.get("/reports/aged-receivables")
def aged_receivables(db: Session = Depends(get_db), user: User = Depends(READ)):
    return finreports.aged_receivables(db)


@router.get("/reports/aged-payables")
def aged_payables(db: Session = Depends(get_db), user: User = Depends(READ)):
    return finreports.aged_payables(db)


@router.get("/reports/sales-by-customer")
def sales_by_customer(start: str | None = None, end: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    today = datetime.date.today()
    return finreports.sales_by_customer(db, _d(start, "Start", _fy_start(today)), _d(end, "End", today.isoformat()))


@router.get("/reports/sales-by-item")
def sales_by_item(start: str | None = None, end: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    today = datetime.date.today()
    return finreports.sales_by_item(db, _d(start, "Start", _fy_start(today)), _d(end, "End", today.isoformat()))


@router.get("/reports/gst")
def gst(start: str | None = None, end: str | None = None, db: Session = Depends(get_db), user: User = Depends(READ)):
    today = datetime.date.today()
    return finreports.gst_report(db, _d(start, "Start", today.replace(day=1).isoformat()), _d(end, "End", today.isoformat()))


@router.get("/reports/gst/export")
def gst_export(start: str, end: str, db: Session = Depends(get_db), user: User = Depends(READ)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rep = finreports.gst_report(db, _d(start, "Start", ""), _d(end, "End", ""))
    wb = Workbook()

    def sheet(title, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))

    sheet("b2b", ["GSTIN", "Customer", "Invoice no", "Date", "Place of supply", "Invoice value", "Rate %", "Taxable", "IGST", "CGST", "SGST"],
          [[r["gstin"], r["customer"], r["invoice_no"], r["date"], r["state_code"], r["value"], r["rate"], r["taxable"], r["igst"], r["cgst"], r["sgst"]] for r in rep["b2b"]], first=True)
    sheet("b2c", ["Place of supply", "Rate %", "Taxable", "IGST", "CGST", "SGST"],
          [[r["state_code"], r["rate"], r["taxable"], r["igst"], r["cgst"], r["sgst"]] for r in rep["b2c"]])
    sheet("credit-debit notes", ["Type", "GSTIN", "Customer", "Note no", "Date", "Note value", "Rate %", "Taxable", "IGST", "CGST", "SGST"],
          [[r["note_type"], r["gstin"], r["customer"], r["invoice_no"], r["date"], r["value"], r["rate"], r["taxable"], r["igst"], r["cgst"], r["sgst"]] for r in rep["cdn"]])
    sheet("hsn", ["HSN/SAC", "Rate %", "Quantity", "Taxable value"], [[r["hsn_sac"], r["rate"], r["qty"], r["taxable"]] for r in rep["hsn"]])
    s = rep["summary"]
    sheet("3B summary", ["Head", "IGST", "CGST", "SGST"],
          [["Outward taxable value", s["outward"]["taxable"], "", ""], ["Output tax", s["outward"]["igst"], s["outward"]["cgst"], s["outward"]["sgst"]],
           ["Input tax credit", s["itc"]["igst"], s["itc"]["cgst"], s["itc"]["sgst"]], ["Net tax payable", s["net_payable"]["igst"], s["net_payable"]["cgst"], s["net_payable"]["sgst"]]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=gst_{start}_to_{end}.xlsx"})


@router.get("/integrity")
def integrity(db: Session = Depends(get_db), user: User = Depends(READ)):
    """Cross-checks that prove the books agree with the source documents."""
    today = datetime.date.today().isoformat()
    tb, bs = finreports.trial_balance(db, today), finreports.balance_sheet(db, today)
    ar, ap = finreports.aged_receivables(db), finreports.aged_payables(db)
    checks = [
        {"name": "Trial balance: debits equal credits", "ok": tb["balanced"], "detail": f"₹{tb['total_debit']:,.2f} vs ₹{tb['total_credit']:,.2f}"},
        {"name": "Balance sheet: assets = liabilities + equity", "ok": bs["balanced"], "detail": f"₹{bs['total_assets']:,.2f} vs ₹{bs['total_liabilities'] + bs['total_equity']:,.2f}"},
        {"name": "Receivables ledger agrees with open invoices", "ok": ar["reconciled"], "detail": f"ledger ₹{ar['ledger_balance']:,.2f} vs invoices ₹{ar['totals']['total']:,.2f}"},
        {"name": "Payables ledger agrees with open bills", "ok": ap["reconciled"], "detail": f"ledger ₹{ap['ledger_balance']:,.2f} vs bills ₹{ap['totals']['total']:,.2f}"},
    ]
    return {"ok": all(c["ok"] for c in checks), "checks": checks}
