import datetime
import io
import json

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, defer

from app.database import get_db
from app.models import Agreement, BillingAccount, BillingException, BillingRun, Client, Commercial, CreditNote, Invoice, Payment, User
from app.services import anomalies, billing_v2, csv_parser, invoicing, review, settings_store, zoho_push, snapshots as usage_snapshots, usage_db, zoho
from app.services.auth import require_user, require_write
from app.services.engine import compute_invoice, gst_breakup
from app.routers.core_routes import log_action

router = APIRouter(prefix="/api/billing")

HEADER_FILL = PatternFill(start_color="E8622A", end_color="E8622A", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _next_invoice_no(db: Session, prefix: str) -> str:
    count = db.query(Invoice).count()
    return f"{prefix}{2480 + count}"


def _is_annual_due(client: Client, commercial: Commercial, period: str, confirm_ids: set[str]) -> tuple[bool, str]:
    if client.billing_frequency != "annual":
        return True, ""
    if client.enterprise_id not in confirm_ids:
        return False, "annual billing not confirmed this period"
    if not commercial.first_invoice_date:
        return True, ""  # first invoice ever — always due once confirmed
    anniversary_month = commercial.first_invoice_date[5:7]
    period_month = period[5:7]
    if anniversary_month != period_month:
        return False, f"not due until anniversary month {anniversary_month}"
    return True, ""


class RunExists(Exception):
    pass


def _v2_invoice_result(v2: dict, client, org_state_code: str, tds_rate: float) -> tuple[dict | None, str | None]:
    """Turn a billing-v2 account result into the invoice shape used by runs, or explain why it is not invoiced."""
    if not v2["billable"]:
        return None, v2["life"]["reason"] or "not billable"
    blocking = sorted({e["code"] for e in v2["exceptions"] if e["blocks"]})
    if blocking:
        return None, "withheld: " + ", ".join(blocking)
    sub = float(v2["subtotal"])
    if sub <= 0:
        return None, "nothing to bill (no usage and no fixed commitment)"
    gb = gst_breakup(sub, client.state_code, org_state_code, (v2["contract"].gst_rate or 0.18) * 100)
    tds = round(sub * tds_rate / 100.0, 2)
    lines = []
    for l in v2["lines"]:
        if l["amount"] == 0 and "covered_by_subscription" not in l.get("flags", []):
            continue
        lines.append({"label": l["label"], "metric": l["component_type"], "qty": float(l["qty"]) if l.get("qty") is not None else 1,
                      "calc_detail": l.get("detail", ""), "amount": float(l["amount"]), "component_type": l["component_type"],
                      "trace": {"component_type": l["component_type"], "params": l.get("params", {}), "detail": l.get("detail", ""),
                                "contract_version": v2["contract"].version, "usage": {k: float(x) for k, x in v2["usage"].items() if x is not None},
                                "flags": l.get("flags", [])}})
    return {"lines": lines, "subtotal": sub, "gst_type": gb["gst_type"], "gst": gb["gst_total"], "tds": tds,
            "total": round(sub + gb["gst_total"] - tds, 2)}, None


def _generate_run(db: Session, user, period: str, tds_rate: float, confirm_ids: set[str],
                  usage_by_client: dict, unknown_ids: list[str], source: str,
                  force: bool = False, trigger: str = "manual", snaps: dict | None = None) -> dict:
    """Shared by the database pull, the CSV fallback and the scheduler: turn
    per-client usage into draft invoices using each client's locked terms, run
    the anomaly checks, then route the run through the approval policy."""
    existing = db.query(BillingRun).filter(BillingRun.period == period, BillingRun.status != "rejected",
                                           or_(BillingRun.trigger.is_(None), BillingRun.trigger != "adhoc")).first()
    if existing and not force:
        raise RunExists(f"A billing run for {period} already exists (run #{existing.id}, {existing.status}).")
    clients = db.query(Client).all()
    automation = settings_store.get_automation(db)

    settings = settings_store.get_all_settings(db)
    org_state_code = settings["ORG_STATE_CODE"]
    gst_rate = float(settings["DEFAULT_GST_RATE"])
    prefix = settings["INVOICE_PREFIX"]
    default_terms = int(settings.get("PAYMENT_TERMS_DAYS", 15) or 15)

    run = BillingRun(period=period, tds_rate=tds_rate, status="draft", created_by=user.email)
    db.add(run)
    db.commit()
    db.refresh(run)

    v2_out = {"results": {}, "run_exceptions": []}
    if db.query(BillingAccount).first() is not None:
        v2_out = billing_v2.compute_run(db, period, snaps if snaps is not None else usage_snapshots.load(db, period), {c.id: c for c in clients})

    invoices_out = []
    skipped = []
    billed: dict[str, float] = {}
    for client in clients:
        if (client.status or "active") != "active":
            skipped.append({"enterprise_id": client.enterprise_id, "reason": f"customer is {client.status}"})
            continue
        v2 = v2_out["results"].get(client.id)
        commercial = None
        if v2 is not None:
            result, why = _v2_invoice_result(v2, client, org_state_code, tds_rate)
            if why:
                skipped.append({"enterprise_id": client.enterprise_id, "reason": why})
                continue
        else:
            commercial = db.query(Commercial).filter(Commercial.client_id == client.id, Commercial.status == "locked").first()
            if commercial is None:
                skipped.append({"enterprise_id": client.enterprise_id, "reason": "no locked commercial"})
                continue
            due, reason = _is_annual_due(client, commercial, period, confirm_ids)
            if not due:
                skipped.append({"enterprise_id": client.enterprise_id, "reason": reason})
                continue
            usage = usage_by_client.get(client.enterprise_id, {
                "loading": 0, "upi_transactions": 0, "upi_amount": 0, "admin_users": 0, "spender_users": 0, "total_users": 0,
            })
            result = compute_invoice(commercial.rules or [], usage, client.state_code, org_state_code, gst_rate, tds_rate, is_annual_due=True)
            if not result["lines"] or result["subtotal"] <= 0:
                skipped.append({"enterprise_id": client.enterprise_id, "reason": "zero usage / no billable lines"})
                continue


        invoice_no = _next_invoice_no(db, prefix)
        terms_days = int(client.payment_terms_days or default_terms)
        due_date = (datetime.date.today() + datetime.timedelta(days=terms_days)).isoformat()
        invoice = Invoice(
            run_id=run.id, client_id=client.id, invoice_no=invoice_no, period=period,
            lines=result["lines"], subtotal=result["subtotal"], gst_type=result["gst_type"],
            gst=result["gst"], tds=result["tds"], total=result["total"],
            payment_terms_days=terms_days, due_date=due_date, balance_due=result["total"],
        )
        db.add(invoice)
        db.flush()  # keep invoice numbering unique within the run

        if commercial is not None and client.billing_frequency == "annual" and not commercial.first_invoice_date:
            commercial.first_invoice_date = f"{period}-06"

        invoices_out.append(invoice_no)
        billed[client.enterprise_id] = result["subtotal"]

    found = anomalies.detect(db, period, clients, usage_by_client, unknown_ids, billed, skipped,
                             float(automation["ANOMALY_SWING_PCT"] or 50))
    names = {c.id: c for c in clients}
    sev = {"high": "high", "medium": "medium", "low": "low", "info": "low"}
    for cid, r in v2_out["results"].items():
        for e in r["exceptions"]:
            db.add(BillingException(run_id=run.id, account_id=r["account"].id, code=e["code"], severity=e["severity"], message=e["message"], blocks=e["blocks"]))
            if e["severity"] != "info":
                found.append({"severity": sev[e["severity"]], "code": e["code"], "enterprise_id": names[cid].enterprise_id, "message": e["message"]})
    for e in v2_out["run_exceptions"]:
        db.add(BillingException(run_id=run.id, account_id=None, code=e["code"], severity=e["severity"], message=e["message"], blocks=False))
        if e["severity"] != "info":
            found.append({"severity": sev[e["severity"]], "code": e["code"], "enterprise_id": "—", "message": e["message"]})
    order = {"high": 0, "medium": 1, "low": 2}
    found.sort(key=lambda a: (order[a["severity"]], a["enterprise_id"]))
    db.flush()
    queue = review.build(db, run, v2_out["results"], float(automation["ANOMALY_SWING_PCT"] or 50))
    run_total = round(sum(i.total for i in run.invoices) if run.invoices else 0.0, 2)
    run.summary = {
        "invoices_generated": len(invoices_out), "skipped": skipped, "unknown_enterprise_ids": unknown_ids,
        "usage_source": source, "anomalies": found, "total": run_total, "review": queue,
    }
    run.trigger = trigger
    run.submitted_at = datetime.datetime.utcnow()

    # approval policy
    finalized = None
    if not settings_store.is_yes(automation["APPROVAL_REQUIRED"]):
        run.status, run.auto_approved, run.approved_by = "approved", True, "policy: approval not required"
    else:
        limit = float(automation["AUTO_APPROVE_MAX_TOTAL"] or 0)
        if limit > 0 and run_total <= limit and not [a for a in found if a['severity'] != 'low'] and invoices_out and not queue["flagged"]:
            run.status, run.auto_approved, run.approved_by = "approved", True, f"policy: clean run under ₹{limit:,.0f}"
        else:
            run.status = "pending_approval"
    if run.status == "approved":
        run.approved_at = datetime.datetime.utcnow()
        review.auto_clear(db, run.id)
        usage_snapshots.freeze(db, period)
        if settings_store.is_yes(automation["AUTO_ISSUE_ON_APPROVAL"]):
            db.flush()
            finalized = invoicing.finalize_run(db, run)
    db.commit()
    log_action(db, user, "billing.calculate", "billing_run",
               f"run={run.id} invoices={len(invoices_out)} source={source} status={run.status} anomalies={len(found)}")
    return {"run_id": run.id, "invoices": invoices_out, "skipped": skipped, "unknown_enterprise_ids": unknown_ids,
            "usage_source": source, "status": run.status, "anomalies": found, "finalized": finalized}


def run_from_database(db: Session, user, period: str, tds_rate: float, confirm_ids: set[str], force: bool = False, trigger: str = "manual") -> dict:
    """Read the period's usage from the platform database and generate a run. Raises
    usage_db.UsageDbError / RunExists for the caller to translate."""
    loading, upi, users = usage_db.fetch_usage(db, period)
    known_ids = {c.enterprise_id for c in db.query(Client).all()}
    usage_by_client, unknown_ids = csv_parser.aggregate_usage(loading, upi, users, known_ids)
    if db.query(BillingAccount).first() is not None:  # v2 accounts read frozen per-enterprise snapshots
        usage_snapshots.save(db, period, usage_db.fetch_snapshot_rows(db, period), "database")
    return _generate_run(db, user, period, tds_rate, confirm_ids, usage_by_client, unknown_ids, "database", force=force, trigger=trigger,
                         snaps=usage_snapshots.load(db, period))


@router.post("/calculate-db")
def calculate_billing_from_db(
    period: str = Form(...),
    tds_rate: float = Form(0.0),
    confirm_annual: str = Form(""),
    force: bool = Form(False),
    db: Session = Depends(get_db),
    user: User = Depends(require_write("billing")),
):
    """Primary path: read the period's usage directly from the platform database."""
    confirm_ids = {x.strip() for x in confirm_annual.split(",") if x.strip()}
    try:
        return run_from_database(db, user, period, tds_rate, confirm_ids, force)
    except usage_db.UsageDbError as exc:
        raise HTTPException(400, f"Usage database: {exc}")
    except RunExists as exc:
        raise HTTPException(409, str(exc))


@router.post("/calculate")
async def calculate_billing(
    period: str = Form(...),
    tds_rate: float = Form(0.0),
    confirm_annual: str = Form(""),  # comma-separated EnterpriseIds
    force: bool = Form(False),
    loading_file: UploadFile = File(...),
    upi_file: UploadFile = File(...),
    users_file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_write("billing")),
):
    """Fallback path: monthly CSV exports (kept for backfills / when the DB feed is down)."""
    loading = csv_parser.parse_loading_csv(await loading_file.read())
    upi = csv_parser.parse_upi_csv(await upi_file.read())
    users = csv_parser.parse_users_csv(await users_file.read())
    known_ids = {c.enterprise_id for c in db.query(Client).all()}
    usage_by_client, unknown_ids = csv_parser.aggregate_usage(loading, upi, users, known_ids)
    confirm_ids = {x.strip() for x in confirm_annual.split(",") if x.strip()}
    usage_snapshots.save(db, period, usage_snapshots.from_csv(loading, upi, users), "csv")
    try:
        return _generate_run(db, user, period, tds_rate, confirm_ids, usage_by_client, unknown_ids, "csv", force=force,
                             snaps=usage_snapshots.load(db, period))
    except RunExists as exc:
        raise HTTPException(409, str(exc))


@router.get("/runs")
def list_runs(db: Session = Depends(get_db), user: User = Depends(require_user)):
    counts = dict(db.query(Invoice.run_id, func.count(Invoice.id)).group_by(Invoice.run_id).all())
    rows = db.query(BillingRun).order_by(BillingRun.id.desc()).all()
    return [{
        "id": r.id, "period": r.period, "tds_rate": r.tds_rate, "status": r.status,
        "summary": r.summary, "created_by": r.created_by, "created_at": r.created_at.isoformat(),
        "invoice_count": counts.get(r.id, 0), "trigger": r.trigger, "total": (r.summary or {}).get("total", 0),
        "anomaly_count": len((r.summary or {}).get("anomalies", [])),
    } for r in rows]


@router.get("/runs/{run_id}")
def run_detail(run_id: int, db: Session = Depends(get_db), user: User = Depends(require_user)):
    run = db.query(BillingRun).get(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    invoices = []
    for inv in run.invoices:
        client = db.query(Client).get(inv.client_id)
        invoices.append({
            "id": inv.id, "invoice_no": inv.invoice_no, "client_name": client.name if client else "?",
            "enterprise_id": client.enterprise_id if client else "?", "lines": inv.lines,
            "subtotal": inv.subtotal, "gst_type": inv.gst_type, "gst": inv.gst, "tds": inv.tds,
            "total": inv.total, "status": inv.status, "zoho_invoice_id": inv.zoho_invoice_id,
        })
    return {"id": run.id, "period": run.period, "tds_rate": run.tds_rate, "status": run.status,
            "summary": run.summary, "invoices": invoices, "created_by": run.created_by, "trigger": run.trigger,
            "approved_by": run.approved_by, "approved_at": run.approved_at.isoformat() if run.approved_at else None,
            "auto_approved": run.auto_approved, "rejection_reason": run.rejection_reason, "review": review.counts(db, run.id)}


class ApproveIn(BaseModel):
    acknowledge_high: bool = False


@router.post("/runs/{run_id}/approve")
def approve_run(run_id: int, payload: ApproveIn, db: Session = Depends(get_db), user: User = Depends(require_write("approvals"))):
    run = db.query(BillingRun).get(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status != "pending_approval":
        raise HTTPException(400, f"Run is {run.status}, not awaiting approval")
    automation = settings_store.get_automation(db)
    if run.created_by == user.email and not settings_store.is_yes(automation["APPROVAL_ALLOW_SELF"]):
        raise HTTPException(403, "Maker-checker: you generated this run, so someone else in Finance must approve it")
    pending = review.counts(db, run_id)["open"]
    if pending:
        raise HTTPException(409, f"{pending} invoice(s) in the review queue still need a check before this run can be locked")
    high = [a for a in (run.summary or {}).get("anomalies", []) if a["severity"] == "high"]
    if high and not payload.acknowledge_high:
        raise HTTPException(409, f"{len(high)} high-severity finding(s) must be acknowledged before approving")
    run.status, run.approved_by, run.approved_at = "approved", user.email, datetime.datetime.utcnow()
    usage_snapshots.freeze(db, run.period)
    finalized = None
    if settings_store.is_yes(automation["AUTO_ISSUE_ON_APPROVAL"]):
        finalized = invoicing.finalize_run(db, run)
    db.commit()
    log_action(db, user, "billing.approve", "billing_run", f"run={run_id} high_ack={bool(high)} finalized={finalized}")
    return {"ok": True, "status": run.status, "finalized": finalized}


class RejectIn(BaseModel):
    reason: str


@router.post("/runs/{run_id}/reject")
def reject_run(run_id: int, payload: RejectIn, db: Session = Depends(get_db), user: User = Depends(require_write("approvals"))):
    run = db.query(BillingRun).get(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status != "pending_approval":
        raise HTTPException(400, f"Run is {run.status}, not awaiting approval")
    if not payload.reason.strip():
        raise HTTPException(400, "A rejection reason is required")
    run.status, run.rejection_reason, run.approved_by, run.approved_at = "rejected", payload.reason.strip(), user.email, datetime.datetime.utcnow()
    for inv in run.invoices:
        inv.status = "cancelled"
        inv.balance_due = 0.0
    db.commit()
    log_action(db, user, "billing.reject", "billing_run", f"run={run_id}: {payload.reason}")
    return {"ok": True}


@router.post("/runs/{run_id}/issue-all")
def issue_all(run_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    """Manually issue + send an approved run (used when AUTO_ISSUE_ON_APPROVAL is off)."""
    run = db.query(BillingRun).get(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status not in ("approved", "pushed"):
        raise HTTPException(400, "Only an approved run can be issued")
    result = invoicing.finalize_run(db, run)
    log_action(db, user, "billing.issue_all", "billing_run", f"run={run_id} {result}")
    return result


@router.post("/runs/{run_id}/push-zoho")
def push_run_to_zoho(run_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    """Synchronous push (kept for scripts and small runs). The Close screen uses the background job instead."""
    run = db.query(BillingRun).get(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status not in ("approved", "pushed"):
        raise HTTPException(400, "Only an approved run can be pushed to Zoho")
    pushed, failed = [], []
    for inv in zoho_push._pushable(run):
        try:
            zoho_push.push_one(db, run, inv)
            pushed.append(inv.invoice_no)
        except Exception as exc:
            db.rollback()
            inv = db.query(Invoice).get(inv.id)
            inv.zoho_error = str(exc)[:300]
            failed.append({"invoice_no": inv.invoice_no, "reason": str(exc)})
        db.commit()
    if pushed and not failed:
        run.status = "pushed"
        db.commit()
    log_action(db, user, "billing.push_zoho", "billing_run", f"run={run_id} pushed={len(pushed)} failed={len(failed)}")
    return {"pushed": pushed, "failed": failed}


@router.get("/invoices")
def list_invoices(
    response: Response,
    status: str | None = None,
    period: str | None = None,
    q: str | None = None,
    limit: int | None = Query(None, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    user: User = Depends(require_user),
):
    """All invoices (newest first). Optional filters and limit/offset let big books page instead of
    shipping everything; X-Total-Count reports the size of the filtered set."""
    query = db.query(Invoice, Client.name, Client.enterprise_id, BillingRun.status).outerjoin(Client, Client.id == Invoice.client_id).outerjoin(BillingRun, BillingRun.id == Invoice.run_id).options(defer(Invoice.lines))
    if status:
        query = query.filter(Invoice.status == status)
    if period:
        query = query.filter(Invoice.period == period)
    if q and q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(Invoice.invoice_no.ilike(like), Client.name.ilike(like), Client.enterprise_id.ilike(like)))
    response.headers["X-Total-Count"] = str(query.count())
    query = query.order_by(Invoice.id.desc())
    if limit:
        query = query.offset(offset).limit(limit)
    return [{
        "id": inv.id, "invoice_no": inv.invoice_no, "period": inv.period,
        "client_name": name or "?", "enterprise_id": ent or "?",
        "subtotal": inv.subtotal, "gst_type": inv.gst_type, "gst": inv.gst, "tds": inv.tds,
        "total": inv.total, "status": inv.status, "zoho_invoice_id": inv.zoho_invoice_id,
        "due_date": inv.due_date, "paid_amount": inv.paid_amount, "balance_due": inv.balance_due,
        "disputed": inv.disputed, "promise_date": inv.promise_date, "promise_note": inv.promise_note,
        "public_token": inv.public_token, "last_viewed_at": inv.last_viewed_at.isoformat() if inv.last_viewed_at else None,
        "run_status": run_status,
    } for inv, name, ent, run_status in query.all()]


# --------------------------------------------------------------------------
# Native invoice lifecycle — issue / send / PDF / dispute / credit note.
# These never depend on Zoho: an invoice is a real, closable receivable
# whether or not the org also pushes it to Zoho for delivery.
# --------------------------------------------------------------------------

def _invoice_or_404(db: Session, invoice_id: int) -> Invoice:
    inv = db.query(Invoice).get(invoice_id)
    if not inv:
        raise HTTPException(404, "Invoice not found")
    return inv


@router.post("/invoices/{invoice_id}/issue")
def issue_invoice(invoice_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    inv = _invoice_or_404(db, invoice_id)
    if inv.status != "draft":
        raise HTTPException(400, f"Invoice already {inv.status}")
    if inv.run is not None and inv.run.status not in ("approved", "pushed"):
        raise HTTPException(400, f"The billing run is {inv.run.status} — it must be approved before invoices can be issued")
    invoicing.issue(db, inv)
    db.commit()
    log_action(db, user, "invoice.issue", "invoice", inv.invoice_no)
    return {"ok": True, "status": inv.status, "due_date": inv.due_date}


@router.get("/invoices/{invoice_id}/pdf")
def invoice_pdf(invoice_id: int, db: Session = Depends(get_db), user: User = Depends(require_user)):
    from app.services.invoice_pdf import build_invoice_pdf

    inv = _invoice_or_404(db, invoice_id)
    client = db.query(Client).get(inv.client_id)
    org = settings_store.get_all_settings(db)
    pdf_bytes = build_invoice_pdf(org, client, inv)
    return StreamingResponse(io.BytesIO(pdf_bytes), media_type="application/pdf",
                              headers={"Content-Disposition": f"attachment; filename={inv.invoice_no}.pdf"})


@router.post("/invoices/{invoice_id}/send")
def send_invoice(invoice_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    inv = _invoice_or_404(db, invoice_id)
    if inv.status == "draft":
        raise HTTPException(400, "Issue the invoice before sending it")
    emailed = invoicing.send(db, inv)
    db.commit()
    log_action(db, user, "invoice.send", "invoice", f"{inv.invoice_no} emailed={emailed}")
    return {"ok": True, "emailed": emailed}


class DisputeIn(BaseModel):
    reason: str = ""
    disputed: bool = True


@router.put("/invoices/{invoice_id}/dispute")
def dispute_invoice(invoice_id: int, payload: DisputeIn, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    inv = _invoice_or_404(db, invoice_id)
    inv.disputed = payload.disputed
    inv.dispute_reason = payload.reason if payload.disputed else ""
    if payload.disputed:
        inv.status = "disputed"
    db.commit()
    log_action(db, user, "invoice.dispute", "invoice", f"{inv.invoice_no} disputed={payload.disputed}: {payload.reason}")
    return {"ok": True}


class CreditNoteIn(BaseModel):
    amount: float
    reason: str = ""


@router.post("/invoices/{invoice_id}/credit-note")
def credit_note(invoice_id: int, payload: CreditNoteIn, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    inv = _invoice_or_404(db, invoice_id)
    note = CreditNote(invoice_id=inv.id, amount=payload.amount, reason=payload.reason, created_by=user.email)
    db.add(note)
    inv.balance_due = max(0.0, inv.balance_due - payload.amount)
    if inv.balance_due <= 0.01:
        inv.status = "written_off" if payload.reason.lower().startswith("write") else "paid"
    db.commit()
    log_action(db, user, "invoice.credit_note", "invoice", f"{inv.invoice_no} amount={payload.amount} reason={payload.reason}")
    return {"ok": True, "balance_due": inv.balance_due, "status": inv.status}


@router.get("/invoices/{invoice_id}/credit-notes")
def list_credit_notes(invoice_id: int, db: Session = Depends(get_db), user: User = Depends(require_user)):
    rows = db.query(CreditNote).filter(CreditNote.invoice_id == invoice_id).all()
    return [{"id": n.id, "amount": n.amount, "reason": n.reason, "created_by": n.created_by,
             "created_at": n.created_at.isoformat()} for n in rows]


@router.post("/reminders/run")
def trigger_reminders(db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    from app.services.reminders import run_reminders
    result = run_reminders(db)
    log_action(db, user, "reminders.run", "invoice", f"sent={len(result['sent'])} logged={len(result['logged_only'])}")
    return result


@router.get("/reminders/log")
def reminders_log(db: Session = Depends(get_db), user: User = Depends(require_user)):
    from app.models import ReminderLog
    rows = db.query(ReminderLog).order_by(ReminderLog.sent_at.desc()).limit(200).all()
    out = []
    for r in rows:
        inv = db.query(Invoice).get(r.invoice_id)
        out.append({"invoice_no": inv.invoice_no if inv else "?", "days_overdue": r.days_overdue,
                     "template": r.template, "channel": r.channel, "sent_at": r.sent_at.isoformat()})
    return out


@router.get("/alerts/expiry")
def expiry_alerts(db: Session = Depends(get_db), user: User = Depends(require_user)):
    today = datetime.date.today()
    agreements = db.query(Agreement).filter(Agreement.end_date.isnot(None)).all()
    danger, warn = [], []
    for a in agreements:
        try:
            end = datetime.date.fromisoformat(a.end_date)
        except (ValueError, TypeError):
            continue
        days_left = (end - today).days
        client = db.query(Client).get(a.client_id)
        item = {"client_name": client.name if client else "?", "file_name": a.file_name,
                "end_date": a.end_date, "days_left": days_left}
        if 0 <= days_left <= 30:
            danger.append(item)
        elif 30 < days_left <= 60:
            warn.append(item)
    return {"danger": danger, "warn": warn}


@router.get("/alerts/annual-due")
def annual_due_alerts(period: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    clients = db.query(Client).filter(Client.billing_frequency == "annual").all()
    due = []
    for client in clients:
        commercial = db.query(Commercial).filter(Commercial.client_id == client.id, Commercial.status == "locked").first()
        if not commercial:
            continue
        if not commercial.first_invoice_date:
            due.append({"enterprise_id": client.enterprise_id, "name": client.name, "reason": "first invoice"})
        elif commercial.first_invoice_date[5:7] == period[5:7]:
            due.append({"enterprise_id": client.enterprise_id, "name": client.name, "reason": "anniversary month"})
    return due


@router.post("/zoho/create-contact/{client_id}")
def create_zoho_contact(client_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("billing"))):
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(404, "Client not found")
    result = zoho.create_contact(client.name, client.gst)
    client.zoho_contact_id = result["contact_id"]
    db.commit()
    log_action(db, user, "zoho.create_contact", "client", f"{client.enterprise_id} -> {result['contact_id']}")
    return result


def _autosize_and_style(ws, headers: list[str]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A"].width = max(14, len(header) + 4)


@router.get("/reports/export")
def export_report(kind: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    wb = Workbook()
    ws = wb.active

    if kind == "invoice_register":
        ws.title = "Invoice Register"
        _autosize_and_style(ws, ["Invoice No", "Client", "Period", "Subtotal", "GST", "TDS", "Total", "GST Type", "Status", "Zoho ID"])
        for inv in db.query(Invoice).order_by(Invoice.id).all():
            client = db.query(Client).get(inv.client_id)
            ws.append([inv.invoice_no, client.name if client else "?", inv.period, inv.subtotal, inv.gst,
                       inv.tds, inv.total, inv.gst_type, inv.status, inv.zoho_invoice_id or ""])
    elif kind == "revenue_summary":
        ws.title = "Revenue Summary"
        _autosize_and_style(ws, ["Period", "Invoice Count", "Subtotal", "GST", "Total", "Status"])
        for run in db.query(BillingRun).order_by(BillingRun.id).all():
            subtotal = sum(i.subtotal for i in run.invoices)
            gst = sum(i.gst for i in run.invoices)
            total = sum(i.total for i in run.invoices)
            ws.append([run.period, len(run.invoices), subtotal, gst, total, run.status])
    elif kind == "client_master":
        ws.title = "Client Master"
        _autosize_and_style(ws, ["EnterpriseId", "Name", "GST", "State", "Billing Frequency", "Commercial Status", "Zoho Contact"])
        for client in db.query(Client).all():
            commercial = db.query(Commercial).filter(Commercial.client_id == client.id).order_by(Commercial.id.desc()).first()
            ws.append([client.enterprise_id, client.name, client.gst, client.state, client.billing_frequency,
                       commercial.status if commercial else "none", client.zoho_contact_id or ""])
    else:
        raise HTTPException(400, "kind must be invoice_register|revenue_summary|client_master")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={kind}.xlsx"},
    )
