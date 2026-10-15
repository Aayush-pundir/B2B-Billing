import csv
import datetime
import io
import json

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Agreement, Client, Commercial, CommercialVersion, Invoice, Metric, Addon, Payment, User, AuditLog
from app.services.auth import require_user, require_write
from app.services.engine import escalate_rules
from app.services import settings_store

router = APIRouter(prefix="/api")


def log_action(db: Session, user: User, action: str, entity: str, detail: str = ""):
    db.add(AuditLog(user=user.email, action=action, entity=entity, detail=detail))
    db.commit()


# --------------------------------------------------------------------------
# Clients
# --------------------------------------------------------------------------

class ClientIn(BaseModel):
    enterprise_id: str
    name: str
    gst: str = ""
    pan: str = ""
    state: str = ""
    state_code: str = ""
    billing_contact_email: str = ""
    billing_frequency: str = "monthly"
    status: str = "active"
    customer_type: str = "business"
    legal_name: str = ""
    trade_name: str = ""
    gst_treatment: str = "registered"
    tan: str = ""
    cin: str = ""
    msme_no: str = ""
    address_line1: str = ""
    address_line2: str = ""
    city: str = ""
    pincode: str = ""
    country: str = "India"
    contact_name: str = ""
    contact_phone: str = ""
    cc_emails: str = ""
    currency: str = "INR"
    payment_terms_days: int | None = None
    credit_limit: float = 0.0
    account_manager: str = ""
    segment: str = ""
    notes: str = ""
    reminders_enabled: bool = True


# Master-data columns, in the order used by the bulk template / export.
CLIENT_FIELDS = list(ClientIn.model_fields.keys())
CLIENT_INT_FIELDS = {"payment_terms_days"}
CLIENT_FLOAT_FIELDS = {"credit_limit"}
CLIENT_BOOL_FIELDS = {"reminders_enabled"}
CLIENT_ALIASES = {"EnterpriseId": "enterprise_id", "Name": "name", "GSTIN": "gst", "PAN": "pan", "State": "state", "StateCode": "state_code"}


def _client_dict(c: Client, commercial) -> dict:
    d = {f: getattr(c, f) for f in CLIENT_FIELDS}
    d.update({"id": c.id, "zoho_contact_id": c.zoho_contact_id, "status_reason": c.status_reason or "", "commercial_status": commercial.status if commercial else "none"})
    return d


def _coerce_client_row(row: dict) -> dict:
    """Normalise one CSV/XLSX row into ClientIn-shaped kwargs (blank cells ignored)."""
    out: dict = {}
    for raw_key, raw_val in row.items():
        key = CLIENT_ALIASES.get(str(raw_key or "").strip(), str(raw_key or "").strip())
        if key not in CLIENT_FIELDS or raw_val is None or str(raw_val).strip() == "":
            continue
        val = str(raw_val).strip()
        try:
            if key in CLIENT_INT_FIELDS:
                val = int(float(val))
            elif key in CLIENT_FLOAT_FIELDS:
                val = float(val)
            elif key in CLIENT_BOOL_FIELDS:
                val = val.lower() in ("yes", "true", "1", "y")
        except ValueError:
            raise ValueError(f"{key}: '{val}' is not a number")
        out[key] = val
    return out


def read_table(raw: bytes, filename: str) -> list[dict]:
    """Read a CSV or XLSX upload into a list of row dicts (first row = headers)."""
    if (filename or "").lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig", errors="replace")))
    return [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]


@router.get("/clients")
def list_clients(db: Session = Depends(get_db), user: User = Depends(require_user)):
    latest: dict[int, Commercial] = {}
    for com in db.query(Commercial).order_by(Commercial.id).all():  # ascending, so the last write per client is the newest
        latest[com.client_id] = com
    return [_client_dict(c, latest.get(c.id)) for c in db.query(Client).order_by(Client.name).all()]


@router.post("/clients")
def create_client(payload: ClientIn, db: Session = Depends(get_db), user: User = Depends(require_write("clients"))):
    if db.query(Client).filter(Client.enterprise_id == payload.enterprise_id).first():
        raise HTTPException(400, "EnterpriseId already exists")
    client = Client(**payload.model_dump())
    db.add(client)
    db.commit()
    db.refresh(client)
    log_action(db, user, "client.create", "client", client.enterprise_id)
    return {"id": client.id}


@router.put("/clients/{client_id}")
def update_client(client_id: int, payload: ClientIn, db: Session = Depends(get_db), user: User = Depends(require_write("clients"))):
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(404, "Client not found")
    for k, v in payload.model_dump().items():
        setattr(client, k, v)
    db.commit()
    log_action(db, user, "client.update", "client", client.enterprise_id)
    return {"ok": True}


@router.delete("/clients/{client_id}")
def delete_client(client_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("clients"))):
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(404, "Client not found")
    db.delete(client)
    db.commit()
    log_action(db, user, "client.delete", "client", client.enterprise_id)
    return {"ok": True}


@router.post("/clients/bulk")
async def bulk_import_clients(file: UploadFile = File(...), update_existing: bool = False, db: Session = Depends(get_db), user: User = Depends(require_write("clients"))):
    """Bulk create (and optionally update) customer master data from CSV or XLSX.
    Any column from the template may be present; blank cells are left untouched."""
    rows = read_table(await file.read(), file.filename or "")
    created, updated, skipped = 0, 0, []
    for i, row in enumerate(rows, start=2):
        try:
            data = _coerce_client_row(row)
        except ValueError as exc:
            skipped.append(f"row {i}: {exc}")
            continue
        eid = data.get("enterprise_id", "")
        if not eid or not data.get("name", "x"):
            skipped.append(f"row {i}: missing enterprise_id")
            continue
        existing = db.query(Client).filter(Client.enterprise_id == eid).first()
        if existing:
            if not update_existing:
                skipped.append(f"row {i}: {eid} already exists")
                continue
            for k, v in data.items():
                setattr(existing, k, v)
            updated += 1
            continue
        if not data.get("name"):
            skipped.append(f"row {i}: {eid} has no name")
            continue
        db.add(Client(**data))
        created += 1
    db.commit()
    log_action(db, user, "client.bulk_import", "client", f"created={created} updated={updated} skipped={len(skipped)}")
    return {"created": created, "updated": updated, "skipped": skipped}


@router.get("/clients/{client_id}/statement")
def client_statement(client_id: int, db: Session = Depends(get_db), user: User = Depends(require_user)):
    """Running ledger of everything issued to and collected from a client —
    the document finance teams ask for most often."""
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(404, "Client not found")

    invoices = db.query(Invoice).filter(Invoice.client_id == client_id).order_by(Invoice.created_at).all()
    payments = db.query(Payment).filter(Payment.client_id == client_id).order_by(Payment.date).all()

    entries = []
    for inv in invoices:
        entries.append({
            "date": inv.issued_at.isoformat() if inv.issued_at else inv.created_at.isoformat(),
            "type": "invoice", "ref": inv.invoice_no, "debit": inv.total, "credit": 0.0,
            "status": inv.status,
        })
    for p in payments:
        entries.append({
            "date": p.date, "type": "payment", "ref": p.reference or f"payment#{p.id}",
            "debit": 0.0, "credit": p.amount, "status": p.match_confidence or "matched",
        })
    entries.sort(key=lambda e: e["date"])

    running = 0.0
    for e in entries:
        running += e["debit"] - e["credit"]
        e["running_balance"] = round(running, 2)

    return {
        "client": {"id": client.id, "name": client.name, "enterprise_id": client.enterprise_id},
        "entries": entries,
        "total_billed": round(sum(i.total for i in invoices), 2),
        "total_collected": round(sum(p.amount for p in payments), 2),
        "outstanding": round(running, 2),
    }


@router.get("/clients/template")
def client_template():
    header = ",".join(CLIENT_FIELDS) + "\n"
    return StreamingResponse(io.StringIO(header), media_type="text/csv",
                              headers={"Content-Disposition": "attachment; filename=customers_template.csv"})


@router.get("/clients/export")
def client_export(db: Session = Depends(get_db), user: User = Depends(require_user)):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CLIENT_FIELDS)
    for c in db.query(Client).order_by(Client.name).all():
        w.writerow([getattr(c, f) if getattr(c, f) is not None else "" for f in CLIENT_FIELDS])
    return StreamingResponse(io.StringIO(buf.getvalue()), media_type="text/csv",
                              headers={"Content-Disposition": "attachment; filename=customers.csv"})


# --------------------------------------------------------------------------
# Commercials
# --------------------------------------------------------------------------

class CommercialIn(BaseModel):
    client_id: int
    billing_frequency: str = "monthly"
    first_invoice_date: str | None = None
    rules: list[dict] = []


@router.get("/commercials")
def list_commercials(db: Session = Depends(get_db), user: User = Depends(require_user)):
    rows = db.query(Commercial).all()
    return [{
        "id": c.id, "client_id": c.client_id, "billing_frequency": c.billing_frequency,
        "first_invoice_date": c.first_invoice_date, "rules": c.rules, "status": c.status,
        "version_no": c.version_no, "source": c.source or "manual", "source_ref": c.source_ref or "",
        "amended_by": c.amended_by or "", "amended_at": c.amended_at.isoformat() if c.amended_at else None,
    } for c in rows]


@router.post("/commercials")
def save_commercial(payload: CommercialIn, db: Session = Depends(get_db), user: User = Depends(require_write("commercials"))):
    commercial = db.query(Commercial).filter(Commercial.client_id == payload.client_id).first()
    if commercial is None:
        commercial = Commercial(client_id=payload.client_id)
        db.add(commercial)
    if commercial.status == "locked":
        raise HTTPException(400, "Commercial is locked. Unlock before editing.")
    commercial.billing_frequency = payload.billing_frequency
    commercial.first_invoice_date = payload.first_invoice_date
    commercial.rules = payload.rules
    commercial.status = "draft"
    commercial.source, commercial.source_ref = "manual", ""
    commercial.amended_by, commercial.amended_at = user.email, datetime.datetime.utcnow()
    db.commit()
    db.refresh(commercial)
    log_action(db, user, "commercial.save", "commercial", str(commercial.id))
    return {"id": commercial.id}


@router.post("/commercials/{commercial_id}/lock")
def lock_commercial(commercial_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("commercials"))):
    commercial = db.query(Commercial).get(commercial_id)
    if not commercial:
        raise HTTPException(404, "Commercial not found")
    commercial.status = "locked"
    commercial.version_no += 1
    db.add(CommercialVersion(
        commercial_id=commercial.id, version_no=commercial.version_no, rules=commercial.rules,
        billing_frequency=commercial.billing_frequency, snapshot_reason="locked", created_by=user.email,
        source=commercial.source or "manual",
    ))
    db.commit()
    log_action(db, user, "commercial.lock", "commercial", str(commercial.id))
    return {"ok": True, "version_no": commercial.version_no}


@router.post("/commercials/{commercial_id}/unlock")
def unlock_commercial(commercial_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("commercials"))):
    commercial = db.query(Commercial).get(commercial_id)
    if not commercial:
        raise HTTPException(404, "Commercial not found")
    commercial.status = "draft"
    db.commit()
    log_action(db, user, "commercial.unlock", "commercial", str(commercial.id))
    return {"ok": True}


@router.get("/commercials/{commercial_id}/versions")
def commercial_versions(commercial_id: int, db: Session = Depends(get_db), user: User = Depends(require_user)):
    rows = db.query(CommercialVersion).filter(CommercialVersion.commercial_id == commercial_id).order_by(CommercialVersion.version_no.desc()).all()
    return [{
        "version_no": v.version_no, "rules": v.rules, "billing_frequency": v.billing_frequency,
        "snapshot_reason": v.snapshot_reason, "created_by": v.created_by, "created_at": v.created_at.isoformat(),
    } for v in rows]


class EscalateIn(BaseModel):
    pct: float


@router.post("/commercials/{commercial_id}/escalate")
def escalate_commercial(commercial_id: int, payload: EscalateIn, db: Session = Depends(get_db), user: User = Depends(require_write("commercials"))):
    commercial = db.query(Commercial).get(commercial_id)
    if not commercial:
        raise HTTPException(404, "Commercial not found")
    commercial.rules = escalate_rules(commercial.rules or [], payload.pct)
    commercial.status = "draft"
    commercial.version_no += 1
    db.add(CommercialVersion(
        commercial_id=commercial.id, version_no=commercial.version_no, rules=commercial.rules,
        billing_frequency=commercial.billing_frequency, snapshot_reason=f"escalation +{payload.pct}%", created_by=user.email,
    ))
    db.commit()
    log_action(db, user, "commercial.escalate", "commercial", f"{commercial.id} +{payload.pct}%")
    return {"ok": True, "rules": commercial.rules}


# --------------------------------------------------------------------------
# Terms import routes — finance only. Three ways in: manual (POST /commercials),
# from a reviewed agreement extraction, and bulk Excel.
# --------------------------------------------------------------------------

RATE_TYPES = {"percent", "flat", "per_unit", "hybrid", "annual_per_unit", "addon"}
SLAB_MODES = {"cliff", "tiered", "none"}
REQUIRED_BY_TYPE = {
    "percent": [],  # needs rate_pct, or tiers when slab_mode is tiered (checked below)
    "flat": ["flat_amount"],
    "per_unit": ["unit_rate"],
    "hybrid": ["hybrid_fixed_amount"],
    "annual_per_unit": ["unit_rate"],
    "addon": [],
}
RULE_NUMERIC = ["rate_pct", "floor_amount", "flat_amount", "unit_rate", "minimum_qty",
                "hybrid_fixed_amount", "hybrid_included_qty", "hybrid_over_rate", "escalation_pct_annual"]
EXCEL_COLUMNS = ["enterprise_id", "billing_frequency", "first_invoice_date", "label", "metric", "rate_type", "slab_mode",
                 *RULE_NUMERIC, "tiers", "invoice_as"]


def _valid_metric_codes(db: Session) -> set[str]:
    return {m.code for m in db.query(Metric).all()} | {a.code for a in db.query(Addon).all()} | {"trucks"}


def validate_rule(rule: dict, valid_metrics: set[str]) -> list[str]:
    errs = []
    rt = rule.get("rate_type")
    if rt not in RATE_TYPES:
        errs.append(f"rate_type must be one of {sorted(RATE_TYPES)}")
        return errs
    if not rule.get("metric"):
        errs.append("metric is required")
    elif rule["metric"] not in valid_metrics:
        errs.append(f"unknown metric '{rule['metric']}' (known: {sorted(valid_metrics)})")
    if rule.get("slab_mode", "cliff") not in SLAB_MODES:
        errs.append(f"slab_mode must be one of {sorted(SLAB_MODES)}")
    for f in REQUIRED_BY_TYPE[rt]:
        if rule.get(f) in (None, ""):
            errs.append(f"{f} is required for rate_type '{rt}'")
    if rt == "percent":
        if rule.get("slab_mode") == "tiered":
            if not rule.get("tiers"):
                errs.append("tiers are required when slab_mode is 'tiered'")
        elif rule.get("rate_pct") in (None, ""):
            errs.append("rate_pct is required for a percent rule")
    for t in rule.get("tiers") or []:
        if "min" not in t or "rate_pct" not in t:
            errs.append("each tier needs at least min and rate_pct")
            break
    return errs


def _rule_from_excel_row(row: dict) -> dict:
    rule: dict = {}
    for key in ("label", "metric", "rate_type", "slab_mode", "invoice_as"):
        v = row.get(key)
        if v not in (None, ""):
            rule[key] = str(v).strip()
    for key in RULE_NUMERIC:
        v = row.get(key)
        if v not in (None, ""):
            try:
                rule[key] = float(v)
            except ValueError:
                raise ValueError(f"{key}: '{v}' is not a number")
    t = row.get("tiers")
    if t not in (None, ""):
        try:
            rule["tiers"] = json.loads(t) if isinstance(t, str) else t
        except json.JSONDecodeError:
            raise ValueError("tiers must be valid JSON, e.g. [{\"min\":0,\"max\":500000,\"rate_pct\":1.5,\"floor\":5000}]")
    if rule.get("rate_type") == "percent":
        rule.setdefault("slab_mode", "cliff")
    return rule


@router.get("/commercials/excel-template")
def commercials_excel_template(db: Session = Depends(get_db), user: User = Depends(require_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Terms"
    ws.append(EXCEL_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    first = db.query(Client).order_by(Client.id).first()
    eid = first.enterprise_id if first else "ENT-1001"
    ws.append([eid, "monthly", "", "Loading fee", "loading", "percent", "cliff", 1.5, 5000, "", "", "", "", "", "", "", "", "", ""])
    ws.append([eid, "monthly", "", "Platform fee", "total_users", "flat", "", "", "", 10000, "", "", "", "", "", "", "", "", ""])
    help_ws = wb.create_sheet("How to fill")
    for line in [
        "One row per rule. A client with several rules has several rows with the same enterprise_id.",
        "enterprise_id must already exist in Customers. A client's rows are saved together — if any row is invalid, none of that client's rows are saved.",
        f"rate_type: {', '.join(sorted(RATE_TYPES))}",
        f"slab_mode (percent rules): {', '.join(sorted(SLAB_MODES))}",
        "metric: loading, upi_transactions, admin_users, spender_users, total_users, trucks — or an add-on code.",
        "percent+cliff needs rate_pct (and optional floor_amount). percent+tiered needs tiers (JSON). flat needs flat_amount. per_unit / annual_per_unit need unit_rate.",
        "hybrid needs hybrid_fixed_amount, hybrid_included_qty, hybrid_over_rate.",
        "Uploaded terms are saved as DRAFT. Review and lock them on the Commercials page.",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              headers={"Content-Disposition": "attachment; filename=terms_template.xlsx"})


@router.post("/commercials/bulk-excel")
async def commercials_bulk_excel(file: UploadFile = File(...), dry_run: bool = False, db: Session = Depends(get_db),
                                 user: User = Depends(require_write("commercials"))):
    """Bulk-upload terms from Excel/CSV. Validates every row, saves each valid
    client's rules as a DRAFT (never auto-locks). dry_run=true validates only."""
    rows = read_table(await file.read(), file.filename or "")
    if not rows:
        raise HTTPException(400, "The file has no data rows")
    valid_metrics = _valid_metric_codes(db)
    clients = {c.enterprise_id: c for c in db.query(Client).all()}

    per_client: dict[str, dict] = {}
    for i, row in enumerate(rows, start=2):
        eid = str(row.get("enterprise_id") or row.get("EnterpriseId") or "").strip()
        bucket = per_client.setdefault(eid, {"rules": [], "errors": [], "frequency": None, "first_invoice_date": None})
        if not eid:
            bucket["errors"].append(f"row {i}: enterprise_id is blank")
            continue
        if row.get("billing_frequency"):
            bucket["frequency"] = str(row["billing_frequency"]).strip().lower()
        if row.get("first_invoice_date"):
            bucket["first_invoice_date"] = str(row["first_invoice_date"])[:10]
        try:
            rule = _rule_from_excel_row(row)
        except ValueError as exc:
            bucket["errors"].append(f"row {i}: {exc}")
            continue
        errs = validate_rule(rule, valid_metrics)
        if errs:
            bucket["errors"].extend(f"row {i}: {e}" for e in errs)
        else:
            bucket["rules"].append(rule)

    results = []
    for eid, b in per_client.items():
        client = clients.get(eid)
        if client is None:
            b["errors"].append(f"unknown enterprise_id '{eid}' — add the customer first")
        commercial = db.query(Commercial).filter(Commercial.client_id == client.id).first() if client else None
        if commercial and commercial.status == "locked":
            b["errors"].append("terms are locked — unlock before uploading new terms")
        if not b["rules"] and not b["errors"]:
            b["errors"].append("no rules found")
        status = "error" if b["errors"] else ("valid" if dry_run else "saved")
        if status == "saved":
            if commercial is None:
                commercial = Commercial(client_id=client.id)
                db.add(commercial)
            commercial.rules = b["rules"]
            commercial.billing_frequency = b["frequency"] or client.billing_frequency or "monthly"
            if b["first_invoice_date"]:
                commercial.first_invoice_date = b["first_invoice_date"]
            commercial.status = "draft"
            commercial.source, commercial.source_ref = "excel", file.filename or ""
            commercial.amended_by, commercial.amended_at = user.email, datetime.datetime.utcnow()
        results.append({"enterprise_id": eid or "(blank)", "rules": len(b["rules"]), "status": status, "errors": b["errors"]})
    if not dry_run:
        db.commit()
        log_action(db, user, "commercial.bulk_excel", "commercial",
                   f"saved={sum(r['status'] == 'saved' for r in results)} errors={sum(r['status'] == 'error' for r in results)} file={file.filename}")
    return {"dry_run": dry_run, "results": results,
            "saved": sum(r["status"] == "saved" for r in results), "failed": sum(r["status"] == "error" for r in results)}


@router.post("/commercials/from-agreement/{agreement_id}")
def commercial_from_agreement(agreement_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("commercials"))):
    """Import the reviewed extraction of an uploaded agreement as DRAFT terms.
    Extraction alone never touches billing terms; finance imports it explicitly."""
    agreement = db.query(Agreement).get(agreement_id)
    if not agreement:
        raise HTTPException(404, "Agreement not found")
    if not agreement.extracted_json:
        raise HTTPException(400, "Run extraction on this agreement first")
    commercial = db.query(Commercial).filter(Commercial.client_id == agreement.client_id).first()
    if commercial and commercial.status == "locked":
        raise HTTPException(400, "Terms are locked. Unlock before importing new terms.")
    result = agreement.extracted_json
    if commercial is None:
        commercial = Commercial(client_id=agreement.client_id)
        db.add(commercial)
    commercial.rules = result.get("rules", [])
    commercial.status = "draft"
    commercial.source, commercial.source_ref = "agreement", agreement.file_name
    commercial.amended_by, commercial.amended_at = user.email, datetime.datetime.utcnow()
    if agreement.start_date and not commercial.first_invoice_date:
        commercial.first_invoice_date = agreement.start_date
    db.commit()
    db.refresh(commercial)
    log_action(db, user, "commercial.import_agreement", "commercial", f"{commercial.id} from agreement {agreement.id}")
    return {"id": commercial.id, "rules": commercial.rules, "overall_confidence": result.get("overall_confidence")}


# --------------------------------------------------------------------------
# Metrics & Add-ons
# --------------------------------------------------------------------------

class MetricIn(BaseModel):
    code: str
    label: str
    basis: str = "amount"
    active: bool = True


@router.get("/metrics")
def list_metrics(db: Session = Depends(get_db), user: User = Depends(require_user)):
    return [{"id": m.id, "code": m.code, "label": m.label, "basis": m.basis, "active": m.active} for m in db.query(Metric).all()]


@router.post("/metrics")
def create_metric(payload: MetricIn, db: Session = Depends(get_db), user: User = Depends(require_write("metrics"))):
    m = Metric(**payload.model_dump())
    db.add(m)
    db.commit()
    return {"id": m.id}


@router.put("/metrics/{metric_id}")
def update_metric(metric_id: int, payload: MetricIn, db: Session = Depends(get_db), user: User = Depends(require_write("metrics"))):
    m = db.query(Metric).get(metric_id)
    if not m:
        raise HTTPException(404, "Metric not found")
    for k, v in payload.model_dump().items():
        setattr(m, k, v)
    db.commit()
    return {"ok": True}


@router.delete("/metrics/{metric_id}")
def delete_metric(metric_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("metrics"))):
    m = db.query(Metric).get(metric_id)
    if m:
        db.delete(m)
        db.commit()
    return {"ok": True}


class AddonIn(BaseModel):
    code: str
    name: str
    default_rate_type: str = "flat"
    active: bool = True


@router.get("/addons")
def list_addons(db: Session = Depends(get_db), user: User = Depends(require_user)):
    return [{"id": a.id, "code": a.code, "name": a.name, "default_rate_type": a.default_rate_type, "active": a.active} for a in db.query(Addon).all()]


@router.post("/addons")
def create_addon(payload: AddonIn, db: Session = Depends(get_db), user: User = Depends(require_write("addons"))):
    a = Addon(**payload.model_dump())
    db.add(a)
    db.commit()
    return {"id": a.id}


# --------------------------------------------------------------------------
# Users (platform users, not clients)
# --------------------------------------------------------------------------

class UserIn(BaseModel):
    email: str
    role: str = "viewer"


@router.get("/users")
def list_users(db: Session = Depends(get_db), user: User = Depends(require_write("users"))):
    return [{"id": u.id, "email": u.email, "role": u.role, "provider": u.provider,
             "last_login": u.last_login.isoformat() if u.last_login else None} for u in db.query(User).all()]


@router.post("/users")
def create_user(payload: UserIn, db: Session = Depends(get_db), user: User = Depends(require_write("users"))):
    if db.query(User).filter(User.email == payload.email).first():
        raise HTTPException(400, "User already exists")
    u = User(email=payload.email, role=payload.role, provider="dev")
    db.add(u)
    db.commit()
    return {"id": u.id}


@router.put("/users/{user_id}")
def update_user(user_id: int, payload: UserIn, db: Session = Depends(get_db), user: User = Depends(require_write("users"))):
    u = db.query(User).get(user_id)
    if not u:
        raise HTTPException(404, "User not found")
    u.email = payload.email
    u.role = payload.role
    db.commit()
    return {"ok": True}


@router.delete("/users/{user_id}")
def delete_user(user_id: int, db: Session = Depends(get_db), user: User = Depends(require_write("users"))):
    u = db.query(User).get(user_id)
    if u:
        db.delete(u)
        db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Audit
# --------------------------------------------------------------------------

@router.get("/audit")
def audit_log(db: Session = Depends(get_db), user: User = Depends(require_user)):
    rows = db.query(AuditLog).order_by(AuditLog.ts.desc()).limit(200).all()
    return [{"user": r.user, "action": r.action, "entity": r.entity, "detail": r.detail, "ts": r.ts.isoformat()} for r in rows]


# --------------------------------------------------------------------------
# Templates (billing CSV templates pre-filled with active EnterpriseIds)
# --------------------------------------------------------------------------

TEMPLATE_HEADERS = {
    "loading": "TransactionId,EnterpriseId,VirtualAccountNumber,Amount,Fee,TotalTransactionAmount,TransactionType,TransactionDate,Status,Vpan,TransactionCategory,ReferenceTransactionId,OmniTransactionId,PreviousStatus,OriginalTransactionDate,UpdatedOn,RunningBalance,PrevBalance,UtrNo,AddMoneyRequestId,EnterpriseName",
    "upi": "TransactionId,TransactionDate,CustomerId,Amount,TotalTrxnAmount,TransactionCategory,TransactionType,Status,EnterpriseName,EnterpriseId",
    "users": "S.No,EnterpriseId,EnterpriseName,Admin,CO_ADMIN,CHECKER,APPROVER,MAKER,TotalUsers",
}


@router.get("/templates/{kind}")
def download_template(kind: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    if kind not in TEMPLATE_HEADERS:
        raise HTTPException(404, "Unknown template kind")
    clients = db.query(Client).all()
    lines = [TEMPLATE_HEADERS[kind]]
    if kind == "users":
        for i, c in enumerate(clients, start=1):
            lines.append(f"{i},{c.enterprise_id},{c.name},0,0,0,0,0,0")
    else:
        for c in clients:
            lines.append(c.enterprise_id)
    body = "\n".join(lines) + "\n"
    return StreamingResponse(io.StringIO(body), media_type="text/csv",
                              headers={"Content-Disposition": f"attachment; filename={kind}_template.csv"})


# --------------------------------------------------------------------------
# Settings
# --------------------------------------------------------------------------

@router.get("/settings")
def get_settings_route(db: Session = Depends(get_db), user: User = Depends(require_user)):
    return settings_store.get_all_settings(db)


class SettingsIn(BaseModel):
    values: dict[str, str]


@router.put("/settings")
def put_settings(payload: SettingsIn, db: Session = Depends(get_db), user: User = Depends(require_write("settings"))):
    result = settings_store.update_settings(db, payload.values)
    log_action(db, user, "settings.update", "settings", str(list(payload.values.keys())))
    return result
